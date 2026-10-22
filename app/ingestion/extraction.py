"""Text extraction.

Each extractor returns a list of pages so that page numbers survive into the
chunk table and, from there, into citations. Formats with no page concept
(DOCX, plain text) return a single page with ``page_number=None`` rather than
a fabricated 1 - a citation that claims a page the user cannot find is worse
than one that omits it.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass

from app.core.exceptions import ValidationFailed
from app.ingestion.validation import DOCX, MARKDOWN, PDF, TXT, FileKind

logger = logging.getLogger(__name__)


@dataclass(frozen=True, slots=True)
class ExtractedPage:
    page_number: int | None
    text: str


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    pages: list[ExtractedPage]
    title: str | None = None

    @property
    def page_count(self) -> int | None:
        if len(self.pages) == 1 and self.pages[0].page_number is None:
            return None
        return len(self.pages)

    @property
    def is_empty(self) -> bool:
        return not any(page.text.strip() for page in self.pages)


def _extract_pdf(content: bytes) -> ExtractedDocument:
    import pymupdf

    try:
        document = pymupdf.open(stream=content, filetype="pdf")
    except Exception as exc:  # pymupdf raises a variety of errors on bad input
        raise ValidationFailed("The PDF could not be opened; it may be corrupt.") from exc

    with document:
        if document.needs_pass:
            raise ValidationFailed("The PDF is password-protected and cannot be processed.")
        pages = [
            ExtractedPage(page_number=index + 1, text=page.get_text("text") or "")
            for index, page in enumerate(document)
        ]
        title = (document.metadata or {}).get("title") or None

    return ExtractedDocument(pages=pages, title=title.strip() if title else None)


def _extract_docx(content: bytes) -> ExtractedDocument:
    from io import BytesIO

    import docx

    try:
        document = docx.Document(BytesIO(content))
    except Exception as exc:
        raise ValidationFailed("The DOCX file could not be opened; it may be corrupt.") from exc

    blocks: list[str] = [p.text for p in document.paragraphs]
    # Table text is frequently where the substance lives in policy documents,
    # and python-docx keeps it out of `paragraphs`.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))

    title = (document.core_properties.title or "").strip() or None
    text = "\n\n".join(block for block in blocks if block.strip())
    return ExtractedDocument(pages=[ExtractedPage(page_number=None, text=text)], title=title)


def _extract_text(content: bytes) -> ExtractedDocument:
    # Validation has already proved this decodes as UTF-8.
    text = content.decode("utf-8")
    return ExtractedDocument(pages=[ExtractedPage(page_number=None, text=text)])


_EXTRACTORS = {
    PDF.name: _extract_pdf,
    DOCX.name: _extract_docx,
    TXT.name: _extract_text,
    MARKDOWN.name: _extract_text,
}


def extract_text(content: bytes, kind: FileKind) -> ExtractedDocument:
    """Extract text from a validated upload."""
    extractor = _EXTRACTORS.get(kind.name)
    if extractor is None:  # pragma: no cover - guarded by validation
        raise ValidationFailed(f"No text extractor for {kind.name!r}.")

    document = extractor(content)
    if document.is_empty:
        raise ValidationFailed(
            "No text could be extracted from the document. "
            "Scanned images require OCR, which is not enabled."
        )
    return document

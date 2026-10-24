"""Turn an authored Markdown source into PDF, DOCX or plain-text bytes.

The demo corpus is authored once, in Markdown, and rendered into three
formats so that the ingestion pipeline is exercised across all of them -
including the page-number path, which only PDF has.

Rendering happens at seed time rather than being committed, for two reasons.
A text-only repository is reviewable: a change to a demo document shows up in
a diff as a change to a sentence. And rendering with ``pymupdf`` and
``python-docx`` - the same libraries ``app.ingestion.extraction`` reads with -
means the output is a genuine PDF or OPC package, never a file that merely
has the right extension.

Nothing here touches the database or the network.
"""

from __future__ import annotations

import textwrap
from dataclasses import dataclass
from io import BytesIO

#: A4 in PostScript points.
_PAGE_WIDTH = 595.0
_PAGE_HEIGHT = 842.0
_MARGIN_X = 56.0
_MARGIN_TOP = 72.0
_MARGIN_BOTTOM = 64.0

_BODY_SIZE = 10.0
_LINE_HEIGHT = 13.5


@dataclass(frozen=True, slots=True)
class _Line:
    """One laid-out line of a PDF page."""

    text: str
    font: str
    size: float
    #: Extra vertical space *before* this line, in points.
    lead: float = 0.0


def _wrap(text: str, width: int) -> list[str]:
    """Break a paragraph into lines. Indentation is applied by the caller."""
    if not text:
        return [""]
    return textwrap.wrap(
        text,
        width=width,
        break_long_words=False,
        break_on_hyphens=False,
    ) or [""]


def _layout(markdown: str) -> list[_Line]:
    """Flatten Markdown into a list of positioned lines.

    Deliberately a small subset: headings, bullets, numbered items, table
    rows and paragraphs. The demo sources are written to that subset, and a
    full Markdown implementation here would be a second project.
    """
    lines: list[_Line] = []
    paragraph: list[str] = []

    def flush() -> None:
        if not paragraph:
            return
        joined = " ".join(paragraph)
        for index, part in enumerate(_wrap(joined, 92)):
            lines.append(_Line(part, "helv", _BODY_SIZE, lead=6.0 if index == 0 else 0.0))
        paragraph.clear()

    for raw in markdown.splitlines():
        stripped = raw.strip()

        if not stripped:
            flush()
            continue

        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            heading = stripped[level:].strip()
            size = {1: 17.0, 2: 12.5}.get(level, 11.0)
            for index, part in enumerate(_wrap(heading, 64)):
                lines.append(_Line(part, "hebo", size, lead=16.0 if index == 0 else 0.0))
            continue

        if stripped.startswith("|"):
            # Table rows keep their column alignment only in a fixed-width
            # font, so they get one.
            flush()
            lines.append(_Line(stripped, "cour", 8.0, lead=2.0))
            continue

        if stripped.startswith(("- ", "* ")):
            flush()
            for index, part in enumerate(_wrap(stripped[2:], 86)):
                lines.append(
                    _Line(("  - " if index == 0 else "    ") + part, "helv", _BODY_SIZE, lead=3.0)
                )
            continue

        head, _, rest = stripped.partition(". ")
        if head.isdigit() and rest:
            flush()
            for index, part in enumerate(_wrap(rest, 86)):
                prefix = f"  {head}. " if index == 0 else "     "
                lines.append(_Line(prefix + part, "helv", _BODY_SIZE, lead=3.0))
            continue

        if raw.startswith("    "):
            # An indented block: preformatted, kept verbatim.
            flush()
            lines.append(_Line(raw[2:], "cour", 8.5, lead=2.0))
            continue

        paragraph.append(stripped)

    flush()
    return lines


def render_pdf(markdown: str, *, title: str) -> bytes:
    """Lay the source out as a paginated PDF with a real text layer."""
    import pymupdf

    document = pymupdf.open()
    try:
        document.set_metadata(
            {
                "title": title,
                "author": "Helix Labs (fictional demo data)",
                "subject": "Demo document for the Enterprise RAG Assistant",
            }
        )

        page = document.new_page(width=_PAGE_WIDTH, height=_PAGE_HEIGHT)
        cursor = _MARGIN_TOP

        for line in _layout(markdown):
            advance = line.lead + _LINE_HEIGHT
            if cursor + advance > _PAGE_HEIGHT - _MARGIN_BOTTOM:
                page = document.new_page(width=_PAGE_WIDTH, height=_PAGE_HEIGHT)
                cursor = _MARGIN_TOP
                advance = _LINE_HEIGHT
            cursor += advance
            if line.text:
                page.insert_text(
                    (_MARGIN_X, cursor),
                    line.text,
                    fontname=line.font,
                    fontsize=line.size,
                )

        return document.tobytes()
    finally:
        document.close()


def render_docx(markdown: str, *, title: str) -> bytes:
    """Write the source as a Word document."""
    import docx

    document = docx.Document()
    document.core_properties.title = title
    document.core_properties.author = "Helix Labs (fictional demo data)"

    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            document.add_paragraph(" ".join(paragraph))
            paragraph.clear()

    for raw in markdown.splitlines():
        stripped = raw.strip()

        if not stripped:
            flush()
            continue

        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            document.add_heading(stripped[level:].strip(), level=min(level, 4))
            continue

        if stripped.startswith(("- ", "* ")):
            flush()
            document.add_paragraph(stripped[2:], style="List Bullet")
            continue

        head, _, rest = stripped.partition(". ")
        if head.isdigit() and rest:
            flush()
            document.add_paragraph(rest, style="List Number")
            continue

        if stripped.startswith("|"):
            # Rendered as a plain line rather than a Word table: what matters
            # downstream is that the text is extractable, and a one-cell-per-
            # line table would chunk worse than the row does.
            flush()
            document.add_paragraph(stripped.strip("|").replace("|", " - ").strip())
            continue

        paragraph.append(stripped)

    flush()

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_txt(markdown: str, *, title: str) -> bytes:
    """Plain text: the authored source, unchanged, as UTF-8.

    ``title`` is accepted so the three renderers share one signature; a text
    file has nowhere to put it, and the source already opens with the title
    as its first heading.
    """
    del title
    return markdown.encode("utf-8")


RENDERERS = {"pdf": render_pdf, "docx": render_docx, "txt": render_txt}

CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
}


def render(markdown: str, *, fmt: str, title: str) -> tuple[bytes, str]:
    """Render ``markdown`` to ``fmt``; returns the bytes and the MIME type."""
    try:
        renderer = RENDERERS[fmt]
    except KeyError:
        raise ValueError(f"Unsupported demo document format: {fmt!r}") from None
    return renderer(markdown, title=title), CONTENT_TYPES[fmt]
